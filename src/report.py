"""Write a color-coded XLSX report using openpyxl.

Color coding by conflict status:
  - red   : CONFLICT
  - green : CLEAR
  - amber : UNVERIFIED

The header row is frozen so it stays visible while scrolling. The same
function serves both screen mode and source mode; source mode adds a
"Search Query" column and a "Channel Country" column (the channel's OWN
stated country, NOT audience location).
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src import conflict as conflict_mod

# Fill colors per status (solid).
_FILLS = {
    conflict_mod.CONFLICT: PatternFill("solid", fgColor="F8CBAD"),   # red-ish
    conflict_mod.CLEAR: PatternFill("solid", fgColor="C6EFCE"),      # green
    conflict_mod.UNVERIFIED: PatternFill("solid", fgColor="FFEB9C"), # amber
}

_NOTES_HEADER = "Negotiation Notes"


def _format_subs(r):
    subs = r.get("subscribers")
    return subs if subs is not None else ""


def _format_tier(r):
    tier = r.get("tier_one_pct")
    return f"{tier:.0f}%" if tier is not None else ""


def _format_notes(r):
    return "\n".join("• " + b for b in r.get("bullets", []))


# Column spec: (header, getter, width, wrap). Built once per report so adding
# the source-mode columns can't desync the coloring / freeze / width logic.
def _columns(include_source):
    cols = [
        ("Handle", lambda r: r.get("handle"), 18, False),
        ("Name", lambda r: r.get("name") or "", 22, False),
        ("Platform", lambda r: r.get("platform") or "", 12, False),
        ("Subscribers", _format_subs, 13, False),
        ("Verified", lambda r: "yes" if r.get("verified") else "no", 9, False),
        ("Status", lambda r: r.get("status"), 12, False),
        ("Category", lambda r: r.get("category") or "", 16, False),
        ("Tier-1 %", _format_tier, 9, False),
        ("Leverage", lambda r: r.get("leverage"), 11, False),
        ("Conflict Reason", lambda r: r.get("conflict_reason") or "", 40, False),
        (_NOTES_HEADER, _format_notes, 50, True),
        ("Source", lambda r: r.get("source") or "", 28, False),
    ]
    if include_source:
        # Search query right after the handle.
        cols.insert(1, ("Search Query", lambda r: r.get("search_query") or "", 18, False))
        # Channel's own stated country, clearly distinguished from audience.
        cols.insert(
            5,
            ("Channel Country (stated, not audience)",
             lambda r: r.get("channel_country") or "", 24, False),
        )
    return cols


def _render_sheet(ws, title, cols, rows, fill_for):
    """Shared sheet renderer: dark frozen header, per-row fill, widths.

    ``cols`` is a list of (header, getter, width, wrap). ``fill_for(row)``
    returns a PatternFill (or None) used to color that row.
    """
    ws.title = title

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="404040")
    for col_idx, (header, _getter, _w, _wrap) in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for row in rows:
        row_idx = ws.max_row + 1
        for col_idx, (_header, getter, _w, wrap) in enumerate(cols, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=getter(row))
            if wrap:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        fill = fill_for(row)
        if fill:
            for col_idx in range(1, len(cols) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = fill

    ws.freeze_panes = "A2"
    for col_idx, (_header, _getter, width, _wrap) in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def write_report(results, out_path, include_source_columns=False,
                 needs_verification=None):
    """Write screened results to an XLSX file at out_path.

    Set ``include_source_columns=True`` for source mode to add the Search
    Query and Channel Country columns. ``needs_verification`` (source mode) is
    written to a separate sheet for creators whose subscriber count could not
    be retrieved, so the range filter is never applied to a guessed number.
    """
    cols = _columns(include_source_columns)
    wb = Workbook()
    _render_sheet(
        wb.active,
        "Creator Sourcing" if include_source_columns else "Conflict Screen",
        cols,
        results,
        fill_for=lambda r: _FILLS.get(r.get("status")),
    )
    if needs_verification:
        _render_sheet(
            wb.create_sheet("Needs verification"),
            "Needs verification",
            cols,
            needs_verification,
            fill_for=lambda r: _FILLS[conflict_mod.UNVERIFIED],
        )
    wb.save(out_path)
    return out_path


# Market-map coloring: green when the brand maps to a known category, amber
# when it's unmapped — reusing the existing CLEAR/UNVERIFIED palette.
_MARKET_MAPPED_FILL = _FILLS[conflict_mod.CLEAR]
_MARKET_UNMAPPED_FILL = _FILLS[conflict_mod.UNVERIFIED]

_MARKET_COLUMNS = [
    ("Brand", lambda r: r.get("brand"), 26, False),
    ("Distinct creators", lambda r: r.get("distinct_creators"), 16, False),
    ("Total placements", lambda r: r.get("total_placements"), 16, False),
    ("Most recent placement", lambda r: r.get("most_recent_date") or "", 20, False),
    ("Category", lambda r: r.get("category") or "unmapped", 16, False),
    ("Example creators (up to 3)", lambda r: r.get("example_handles") or "", 40, True),
]


def write_market_report(rows, out_path):
    """Write the ranked brand market-map to an XLSX file at out_path."""
    wb = Workbook()
    _render_sheet(
        wb.active,
        "Market Map",
        _MARKET_COLUMNS,
        rows,
        fill_for=lambda r: (
            _MARKET_UNMAPPED_FILL if r.get("category") in (None, "unmapped")
            else _MARKET_MAPPED_FILL
        ),
    )
    wb.save(out_path)
    return out_path


# Radar (football brief) report: ranked brand table with the brief's columns,
# colored by roster conflict, plus a separate EXCLUDED FOR REVIEW sheet.
_RADAR_COLUMNS = [
    ("Brand", lambda r: r.get("brand"), 22, False),
    ("Reachability", lambda r: r.get("reachability") or "", 16, False),
    ("Contact starting point", lambda r: r.get("contact_start") or "", 46, True),
    ("Brief category", lambda r: r.get("brief_category"), 24, False),
    ("Platform", lambda r: r.get("platforms"), 20, False),
    ("Corroboration", lambda r: r.get("corroboration") or "", 22, False),
    ("Recent sponsorships (proof)", lambda r: r.get("proof") or "", 44, True),
    ("Live vs Evergreen", lambda r: r.get("recency_tag"), 18, False),
    ("Suggested roster fit", lambda r: r.get("suggested_fit") or "none", 30, True),
    ("Roster conflict", lambda r: r.get("roster_conflict"), 26, True),
    ("Most recent date", lambda r: r.get("most_recent_date") or "", 16, False),
    ("Verification", lambda r: r.get("verification"), 18, False),
    ("Distinct creators", lambda r: r.get("distinct_creators"), 16, False),
    ("Total placements", lambda r: r.get("total_placements"), 16, False),
]

_REVIEW_COLUMNS = [
    ("Brand (EXCLUDED FOR REVIEW)", lambda r: r.get("brand"), 32, False),
    ("Reason", lambda r: r.get("reason"), 70, True),
]


def _radar_fill(row):
    rc = (row.get("roster_conflict") or "")
    if rc.startswith("CONFLICT"):
        return _FILLS[conflict_mod.CONFLICT]
    if rc == "CLEAR":
        return _FILLS[conflict_mod.CLEAR]
    if rc.startswith("UNVERIFIED"):
        return _FILLS[conflict_mod.UNVERIFIED]
    return None  # n/a (no roster) -> uncolored


def write_radar_report(rows, out_path, excluded_for_review=None,
                       single_mention=None):
    """Write the football radar: ranked brand sheet, a single-mention sheet for
    unknown brands only one creator ran, and an excluded-for-review sheet."""
    wb = Workbook()
    _render_sheet(wb.active, "Football Radar", _RADAR_COLUMNS, rows,
                  fill_for=_radar_fill)
    if single_mention:
        _render_sheet(wb.create_sheet("Single mention (unverified)"),
                      "Single mention (unverified)", _RADAR_COLUMNS,
                      single_mention, fill_for=lambda r: _FILLS[conflict_mod.UNVERIFIED])
    review = wb.create_sheet("Excluded for review")
    _render_sheet(review, "Excluded for review", _REVIEW_COLUMNS,
                  excluded_for_review or [], fill_for=lambda r: _FILLS[conflict_mod.UNVERIFIED])
    wb.save(out_path)
    return out_path
